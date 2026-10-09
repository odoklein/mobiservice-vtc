from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
DOCX_OUT = OUT_DIR / "comparatif-formules-tarification-vtc.docx"

NAVY = "102A43"
BLUE = "1F6FEB"
GREEN = "1F8A5B"
PALE_GREEN = "E9F7F0"
PALE_BLUE = "EAF2FF"
PALE_GOLD = "FFF5D6"
GOLD = "A66B00"
RED = "B42318"
PALE_RED = "FDECEC"
INK = "172B4D"
MUTED = "60758A"
LIGHT = "F4F7FA"
WHITE = "FFFFFF"
BORDER = "CFD9E3"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_run(run, size=10.5, bold=False, color=INK, italic=False, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = rgb(color)


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_para(doc, text="", size=10.5, color=INK, bold=False, italic=False,
             align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep
    set_run(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p


def add_mixed_para(doc, pieces, before=0, after=6, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    for text, opts in pieces:
        set_run(p.add_run(text), **opts)
    return p


def add_heading(doc, text, level=1, before=None, after=None):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    set_run(p.add_run(text), size={1: 16, 2: 13, 3: 11.5}[level],
            bold=True, color=BLUE if level < 3 else NAVY)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    set_run(p.add_run(label.upper() + "  "), size=9, bold=True, color=accent)
    set_run(p.add_run(text), size=10.5, bold=True, color=NAVY)
    add_para(doc, "", size=1, after=3)


def add_formula(doc, title, formula, note=None):
    table = doc.add_table(rows=2 if note else 1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_fill(cell, NAVY)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_run(p.add_run(title + "\n"), size=9, bold=True, color="A8C7FA")
    set_run(p.add_run(formula), size=11.5, bold=True, color=WHITE, font="Consolas")
    if note:
        cell2 = table.cell(1, 0)
        set_cell_fill(cell2, LIGHT)
        p2 = cell2.paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        set_run(p2.add_run(note), size=9.5, color=MUTED)
    add_para(doc, "", size=1, after=3)


def add_table(doc, headers, rows, widths, alignments=None, header_fill=NAVY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_fill(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(header), size=9, bold=True, color=WHITE)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_fill(cells[i], WHITE if r_idx % 2 == 0 else LIGHT)
            p = cells[i].paragraphs[0]
            p.alignment = (alignments[i] if alignments else WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            set_run(p.add_run(str(value)), size=9.2, color=INK,
                    bold=(i == 0 and len(row) > 2))
    set_table_geometry(table, widths)
    add_para(doc, "", size=1, after=4)
    return table


def add_labeled_item(doc, label, text, color=INK):
    add_mixed_para(doc, [
        (label + " — ", {"size": 10.2, "bold": True, "color": color}),
        (text, {"size": 10.2, "color": INK}),
    ], after=5)


def page_break(doc):
    doc.add_page_break()


def configure_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.72)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for level, size, before, after in ((1, 16, 18, 10), (2, 13, 14, 7), (3, 11.5, 10, 5)):
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(BLUE if level < 3 else NAVY)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(hp.add_run("MOBISERVICE VTC  |  COMPARATIF TARIFAIRE"), size=8.5, bold=True, color=MUTED)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(fp.add_run("Document de travail — 31 juillet 2026"), size=8, color=MUTED)
    return doc


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = configure_doc()

    # Page 1 — masthead / executive summary
    add_para(doc, "NOTE DE COMPARAISON", size=9, bold=True, color=GREEN, after=4)
    add_para(doc, "Formule actuelle de la plateforme\nvs. formule transmise", size=25,
             bold=True, color=NAVY, after=7)
    add_para(doc, "Calcul des transferts VTC — aller simple et aller-retour",
             size=13, color=MUTED, after=14)
    add_mixed_para(doc, [
        ("Objet : ", {"size": 9.5, "bold": True, "color": MUTED}),
        ("vérifier les écarts de calcul et les risques d’implémentation", {"size": 9.5, "color": INK}),
    ], after=2)
    add_mixed_para(doc, [
        ("Périmètre : ", {"size": 9.5, "bold": True, "color": MUTED}),
        ("calculateur public /api/pricing/estimate et enregistrement d’une réservation", {"size": 9.5, "color": INK}),
    ], after=14)

    add_callout(
        doc,
        "Conclusion",
        "La formule transmise est identique au calculateur public pour un aller simple, "
        "et donne bien 139,38 € TTC hors péage pour Cluses → Annecy. Les écarts importants "
        "proviennent surtout de l’implémentation lors de l’enregistrement de la réservation.",
        PALE_GREEN,
        GREEN,
    )

    add_heading(doc, "Résultat en un coup d’œil", 1)
    add_table(
        doc,
        ["Point comparé", "Plateforme affichée", "Formule transmise", "Verdict"],
        [
            ["A/S", "CA aller + TP + CA retour + péage ×1", "Même formule", "Identique"],
            ["Départ Cluses", "CA aller = 0", "CA aller omis car nul", "Identique"],
            ["A/R", "CA aller + 2×TP + CA retour + 2×péage + attente", "Même formule", "Identique"],
            ["Attente", "15 min gratuites, puis 1,20/1,80 €/min", "Même règle", "Identique"],
            ["Sauvegarde", "Recalcul séparé sans péage ni attente", "Non prévu", "Écart critique"],
        ],
        [1800, 3000, 2800, 1760],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_callout(
        doc,
        "Décision recommandée",
        "Conserver la formule transmise comme formule de référence, mais utiliser une seule fonction "
        "de calcul côté serveur pour l’estimation, la réservation, le devis et la facture.",
        PALE_GOLD,
        GOLD,
    )

    # Page 2 — exact formulas
    page_break(doc)
    add_heading(doc, "1. Les deux formules, mises à plat", 1, before=0)
    add_para(doc, "Les distances sont calculées depuis le dépôt fixe : 4 rue des Artisans, 74300 Cluses.",
             size=10, color=MUTED, after=10)

    add_heading(doc, "Formule actuellement affichée par la plateforme", 2)
    add_formula(
        doc,
        "ALLER SIMPLE",
        "Prix TTC = (CA aller × tarif CA) + (TP × tarif TP)\n"
        "          + (CA retour × tarif CA) + péage client",
        "CA retour est toujours inclus. Le péage du trajet client est compté une seule fois.",
    )
    add_formula(
        doc,
        "ALLER-RETOUR",
        "Prix TTC = (CA aller × tarif CA) + (2 × TP × tarif TP)\n"
        "          + (CA retour × tarif CA) + (2 × péage)\n"
        "          + attente facturable",
        "Pour un A/R, CA retour est fixé égal à CA aller, car le client revient au point de prise en charge.",
    )

    add_heading(doc, "Formule transmise", 2)
    add_formula(
        doc,
        "ALLER SIMPLE",
        "Prix total = CA aller + trajet client (TP)\n"
        "           + CA retour + péage client",
        "Chaque composant CA/TP est bien distance × tarif. Péage facturé une seule fois.",
    )
    add_formula(
        doc,
        "PRISE EN CHARGE DIRECTEMENT À CLUSES",
        "Prix = (distance client × tarif TP)\n"
        "     + (distance retour × tarif CA) + péage",
        "Cette écriture ne supprime pas CA aller : elle constate simplement que sa distance est égale à zéro.",
    )

    add_heading(doc, "Constat mathématique", 2)
    add_callout(
        doc,
        "Équivalence",
        "Pour un aller simple, les deux écritures produisent exactement le même résultat lorsque "
        "les mêmes distances, paliers et péages sont utilisés.",
        PALE_GREEN,
        GREEN,
    )

    # Page 3 — rates and example
    page_break(doc)
    add_heading(doc, "2. Tarifs et exemple de contrôle", 1, before=0)
    add_heading(doc, "Grille utilisée dans les deux formules", 2)
    add_table(
        doc,
        ["Élément", "Jour", "Nuit / dimanche / jour férié"],
        [
            ["Trajet client (TP)", "1,32 €/km", "1,90 €/km"],
            ["CA — palier >25–50 km", "1,32 €/km", "1,70 €/km"],
            ["CA — palier >50–75 km", "1,10 €/km", "1,40 €/km"],
            ["CA — palier >75–100 km", "0,90 €/km", "1,10 €/km"],
            ["CA — palier >100 km", "0,70 €/km", "0,70 €/km"],
            ["Forfait ≤25 km", "33,00 € TTC", "47,50 € TTC"],
            ["Attente A/R après 15 min", "1,20 €/min", "1,80 €/min"],
        ],
        [3820, 2200, 3340],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
    )
    add_para(
        doc,
        "Plage jour : 07:00–19:59, hors dimanche et jours fériés français. "
        "Plage nuit : 20:00–06:59, dimanche et jours fériés 24 h/24.",
        size=9.5,
        color=MUTED,
        after=10,
    )

    add_heading(doc, "Exemple commun : Cluses → Annecy", 2)
    add_para(doc, "Hypothèses : prise en charge à Cluses, 69 km, tarif jour, aller simple, hors péage.",
             size=10, color=MUTED, after=7)
    add_table(
        doc,
        ["Composant", "Calcul", "Montant"],
        [
            ["CA aller", "0 km × tarif CA", "0,00 €"],
            ["Trajet client (TP)", "69 km × 1,32 €/km", "91,08 €"],
            ["CA retour", "69 km × 0,70 €/km", "48,30 €"],
            ["Péage", "Hors exemple", "0,00 €"],
            ["TOTAL TTC", "91,08 € + 48,30 €", "139,38 €"],
        ],
        [2500, 4300, 2560],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT],
    )
    add_callout(
        doc,
        "Résultat",
        "Plateforme affichée : 139,38 € TTC. Formule transmise : 139,38 € TTC. Écart : 0,00 €.",
        PALE_GREEN,
        GREEN,
    )
    add_para(
        doc,
        "La formule ne contient ni prise en charge fixe de 10 €, ni prix de trajet à la minute, "
        "ni coefficient de retour à vide, ni supplément automatique J+1.",
        size=10.2,
        bold=True,
        color=NAVY,
        after=0,
    )

    # Page 4 — implementation differences
    page_break(doc)
    add_heading(doc, "3. Différences réelles et risques actuels", 1, before=0)
    add_para(
        doc,
        "Les différences ci-dessous ne viennent pas de la formule commerciale transmise. "
        "Elles viennent des chemins de calcul distincts présents dans la plateforme.",
        size=10.2,
        color=MUTED,
        after=10,
    )
    add_table(
        doc,
        ["Sujet", "Comportement constaté", "Conséquence", "Priorité"],
        [
            ["Enregistrement", "Le serveur recalcule sans péage ni attente.", "Prix enregistré potentiellement inférieur au prix affiché.", "Critique"],
            ["Deux moteurs", "L’estimateur utilise des constantes ; la sauvegarde peut utiliser la base.", "Deux tarifs différents pour une même course.", "Critique"],
            ["Attente affichée", "L’écran annonce 10 min gratuites ; le moteur en déduit 15.", "Information client incohérente.", "Haute"],
            ["Palier A/R", "Le palier CA est choisi avec CA aller + 1×TP + CA retour.", "Le TP retour est facturé mais ne participe pas au choix du palier.", "Haute"],
            ["J+1 à J+3", "L’immobilisation est calculée séparément mais non ajoutée au total.", "Le supplément peut manquer au devis.", "Haute"],
            ["TVA sauvegardée", "La réservation stocke un taux global de 10 %.", "Le détail à 20 % des péages/MAD peut être perdu.", "Moyenne"],
        ],
        [1700, 3300, 3000, 1360],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
    )

    add_heading(doc, "Interprétation du palier", 2)
    add_labeled_item(
        doc,
        "Aller simple",
        "la plateforme choisit le tarif CA à partir de CA aller + TP + CA retour. "
        "C’est cohérent avec l’exemple Cluses → Annecy : 0 + 69 + 69 = 138 km, donc palier >100 km.",
    )
    add_labeled_item(
        doc,
        "Aller-retour",
        "le prix facture 2×TP, mais le palier reste déterminé avec un seul TP. "
        "Cette règle doit être confirmée explicitement avant de figer la nouvelle formule.",
    )

    add_callout(
        doc,
        "Point à valider",
        "Pour un A/R, le palier CA doit-il être déterminé avec CA aller + TP + CA retour "
        "ou avec la distance réellement facturée CA aller + 2×TP + CA retour ?",
        PALE_GOLD,
        GOLD,
    )

    # Page 5 — recommendation / sources
    page_break(doc)
    add_heading(doc, "4. Recommandation de mise en œuvre", 1, before=0)
    add_callout(
        doc,
        "Formule cible",
        "Prix TTC = CA aller + TP facturé + CA retour + péage client + attente/MAD applicable. "
        "Aucune prise en charge fixe, aucun prix à la minute pour le trajet et aucun coefficient de retour à vide.",
        PALE_GREEN,
        GREEN,
    )

    add_heading(doc, "Règles à figer dans une seule fonction serveur", 2)
    add_labeled_item(doc, "1. Distances", "figer CA aller, TP et CA retour issus de la même matrice routière.")
    add_labeled_item(doc, "2. Horaire", "choisir jour/nuit selon l’heure locale de prise en charge, dimanche et jours fériés inclus.")
    add_labeled_item(doc, "3. Palier", "documenter l’assiette exacte du palier CA, particulièrement pour les A/R.")
    add_labeled_item(doc, "4. Péages", "facturer ×1 en A/S et ×2 en A/R, uniquement lorsque le client est transporté.")
    add_labeled_item(doc, "5. Attente", "afficher et appliquer 15 minutes gratuites, puis 1,20 €/min jour ou 1,80 €/min nuit.")
    add_labeled_item(doc, "6. TVA", "conserver séparément transport à 10 % et péage/MAD à 20 %.")
    add_labeled_item(doc, "7. Persistance", "enregistrer exactement le prix calculé par la fonction serveur, avec son détail et sa version tarifaire.")

    add_heading(doc, "Verdict final", 2)
    add_para(
        doc,
        "La formule transmise peut servir de formule commerciale de référence : elle correspond au "
        "calculateur public actuel pour les cas décrits. Le chantier nécessaire n’est pas un changement "
        "de formule, mais l’unification et la sécurisation de son exécution dans toute la plateforme.",
        size=11,
        bold=True,
        color=NAVY,
        after=12,
    )

    add_heading(doc, "Sources techniques vérifiées", 2)
    add_para(
        doc,
        "• app/api/pricing/estimate/route.ts — construction des distances, péages et appel du calculateur public\n"
        "• lib/pricing/tariffs-2026.ts — tarifs, forfait, formule A/S-A/R, attente et TVA\n"
        "• app/api/bookings/route.ts — recalcul et enregistrement serveur\n"
        "• lib/pricing.ts — deuxième moteur et chargement des tarifs configurables\n"
        "• app/(public)/reservation/page.tsx — affichage client et immobilisation J+1 à J+3",
        size=8.8,
        color=MUTED,
        after=0,
    )

    doc.core_properties.title = "Comparatif des formules de tarification VTC"
    doc.core_properties.subject = "Formule plateforme actuelle vs formule transmise"
    doc.core_properties.author = "MobiService VTC"
    doc.core_properties.keywords = "VTC, tarification, CA, TP, péage, comparaison"
    doc.save(DOCX_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    build()
