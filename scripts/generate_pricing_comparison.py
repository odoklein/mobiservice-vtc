from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path('deliverables')
DOCX_OUT = OUT_DIR / 'comparatif-formules-tarification-vtc.docx'

BLUE = '2E74B5'
DARK_BLUE = '1F4D78'
NAVY = '0B2545'
PALE_BLUE = 'E8EEF5'
LIGHT_GRAY = 'F2F4F7'
PALE_GREEN = 'EAF5EE'
PALE_AMBER = 'FFF5DB'
PALE_RED = 'FCECEC'
MUTED = '5F6B78'
BLACK = '111827'
WHITE = 'FFFFFF'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(width))
    tc_w.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_ind = tbl_pr.first_child_found_in('w:tblInd')
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), str(indent))
    tbl_ind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_font(run, size=11, color=BLACK, bold=False, italic=False, name='Calibri'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def style_para(p, before=0, after=6, line=1.10, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align


def add_text(doc, text, size=11, color=BLACK, bold=False, italic=False, before=0, after=6, align=None):
    p = doc.add_paragraph()
    style_para(p, before, after, 1.10, align)
    set_font(p.add_run(text), size, color, bold, italic)
    return p


def add_heading(doc, text, level=1):
    settings = {
        1: (16, BLUE, 16, 8),
        2: (13, BLUE, 12, 6),
        3: (12, DARK_BLUE, 8, 4),
    }
    size, color, before, after = settings[level]
    p = doc.add_paragraph()
    style_para(p, before, after, 1.10)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), size, color, True)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    style_para(p, 0, 4, 1.167)
    set_font(p.add_run(text), 10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    style_para(p, 0, 4, 1.167)
    set_font(p.add_run(text), 10.5)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    style_para(p, 2, 2, 1.10)
    set_font(p.add_run(label + '  '), 10.5, NAVY, True)
    set_font(p.add_run(text), 10.5, BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, rows, widths, header_fill=PALE_BLUE, font_size=9.5, aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = 'Table Grid'
    set_table_geometry(table, widths)
    hrow = table.rows[0]
    set_repeat_table_header(hrow)
    for i, value in enumerate(headers):
        cell = hrow.cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        style_para(p, 0, 0, 1.0, WD_ALIGN_PARAGRAPH.CENTER)
        set_font(p.add_run(value), font_size, NAVY, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            alignment = aligns[i] if aligns else WD_ALIGN_PARAGRAPH.LEFT
            style_para(p, 0, 0, 1.0, alignment)
            set_font(p.add_run(str(value)), font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'PAGE'
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def force_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header = section.header
    hp = header.paragraphs[0]
    style_para(hp, 0, 0, 1.0)
    set_font(hp.add_run('MOBISERVICE VTC  |  COMPARATIF TARIFAIRE'), 8.5, MUTED, True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_para(fp, 0, 0, 1.0)
    set_font(fp.add_run('Document de travail  |  Page '), 8, MUTED)
    add_page_field(fp)

    # Cover / executive summary
    add_text(doc, 'NOTE DE DÉCISION', 10, DARK_BLUE, True, after=12)
    title = add_text(doc, 'Comparatif des formules\nde tarification VTC', 26, NAVY, True, after=8)
    title.paragraph_format.line_spacing = 1.0
    add_text(doc, 'Moteur public actuel vs. nouvelle formule proposée', 13, MUTED, False, after=18)
    add_text(doc, 'Périmètre : transferts au départ de Cluses. Montants TTC dans les exemples. Les péages, durées, suppléments et remises restent des variables de la course.', 10, MUTED, False, after=14)

    add_callout(
        doc,
        'À RETENIR',
        'La formule actuelle est segmentée (dépôt, trajet client, retour) et dépend de paliers. La nouvelle formule est plus lisible : prise en charge + distance + temps + retour à vide + frais éventuels.',
        PALE_BLUE,
    )

    add_heading(doc, 'Décision à prendre', 1)
    add_number(doc, 'Choisir une seule formule de référence pour le calculateur public et les réservations.')
    add_number(doc, 'Définir les paramètres manquants de la nouvelle formule : TVA, règle d’attente, péages, A/R client et forfait négocié au-delà de 250 km.')
    add_number(doc, 'Afficher au client un détail identique à celui enregistré dans la réservation et sur le devis.')

    add_heading(doc, 'Synthèse comparative', 1)
    add_table(doc,
              ['Sujet', 'Formule actuelle', 'Nouvelle formule proposée'],
              [
                  ['Lisibilité', 'Moyenne : distances CA/TP et paliers.', 'Élevée : composants tarifaires explicites.'],
                  ['Prévisibilité', 'Variable selon le trajet dépôt/client et le palier.', 'Facile à expliquer si durée et péages sont connus.'],
                  ['Retour à vide', 'Calculé en km au tarif CA.', 'Calculé par pourcentage de la distance aller.'],
                  ['Temps de trajet', 'Non facturé dans le transfert.', 'Facturé à la minute.'],
                  ['Suppléments', 'Attente A/R après 15 min seulement.', 'Attente après 10 min + J+1/J+2/J+3.'],
              ], [1650, 3855, 3855], font_size=9)

    force_page_break(doc)

    # Existing formula
    add_heading(doc, '1. Formule actuelle du calculateur public', 1)
    add_text(doc, 'La formule ci-dessous décrit le moteur actuellement utilisé par /api/pricing/estimate. Il ne contient ni prise en charge fixe, ni tarif minute, ni supplément J+1/J+2/J+3, ni remise.', 10.5, after=8)

    add_heading(doc, 'Composition', 2)
    add_table(doc,
              ['Élément', 'Aller simple (A/S)', 'Aller-retour (A/R)'],
              [
                  ['Déplacement chauffeur', 'Dépôt → prise en charge', 'Dépôt → prise en charge'],
                  ['Trajet client', 'Prise en charge → destination', 'Prise en charge ↔ destination (×2)'],
                  ['Retour chauffeur', 'Destination → dépôt', 'Prise en charge → dépôt'],
                  ['Péages client', '×1', '×2'],
                  ['Attente', 'Non appliquée', '15 min offertes, puis 1,20 €/min jour ou 1,80 €/min nuit'],
              ], [1900, 3730, 3730], font_size=9)

    add_heading(doc, 'Tarifs actuels', 2)
    add_table(doc,
              ['Période', 'TP (trajet client)', 'CA (déplacement / retour chauffeur)'],
              [
                  ['Jour (7h–20h, hors dimanche/JF)', '1,32 €/km', '1,32 / 1,10 / 0,90 / 0,70 €/km selon palier'],
                  ['Nuit (20h–7h, dimanche/JF)', '1,90 €/km', '1,70 / 1,40 / 1,10 / 0,70 €/km selon palier'],
                  ['Agglomération', 'Forfait ≤ 25 km : 33,00 € jour / 47,50 € nuit', 'Le forfait remplace le calcul transport'],
              ], [1900, 2850, 4610], font_size=9)

    add_callout(doc, 'POINT DE VIGILANCE', 'En A/R, le trajet client est bien facturé ×2, mais le moteur utilise actuellement CA aller + TP + CA retour pour choisir le palier et le forfait. Le TP n’est donc pas doublé pour cette décision de palier.', PALE_AMBER)

    add_heading(doc, 'Atouts et contraintes', 2)
    add_bullet(doc, 'Atout : le prix couvre explicitement les déplacements du chauffeur avant et après la course.')
    add_bullet(doc, 'Atout : les péages ne sont facturés que pendant le transport du client.')
    add_bullet(doc, 'Contrainte : la règle est difficile à expliquer sans afficher les trois segments CA/TP/CA.')
    add_bullet(doc, 'Contrainte : le calculateur public et le moteur de réservation ne sont pas parfaitement alignés, notamment sur les péages et la TVA.')

    force_page_break(doc)

    # Proposed formula
    add_heading(doc, '2. Nouvelle formule proposée', 1)
    add_text(doc, 'Cette formule est proposée dans la conversation. Elle n’est pas encore le comportement du calculateur public.', 10.5, DARK_BLUE, True, after=10)

    add_callout(doc, 'FORMULE', 'Prix = prise en charge + distance aller + temps de trajet + péages aller-retour + retour à vide + suppléments − remise.', PALE_GREEN)

    add_heading(doc, 'Paramètres fournis', 2)
    add_table(doc,
              ['Composant', 'Jour', 'Nuit', 'Règle'],
              [
                  ['Prise en charge', '10,00 €', '10,00 €', 'Montant fixe par course'],
                  ['Distance aller', '1,20 €/km', '1,70 €/km', 'Distance client retenue'],
                  ['Temps de trajet', '0,40 €/min', '0,70 €/min', 'Durée du trajet à définir (réelle ou estimée)'],
                  ['Péages', 'Péage ×2', 'Péage ×2', 'Aller + retour, selon la consigne proposée'],
                  ['Attente', 'Après 10 min', 'Après 10 min', 'Tarif précis à confirmer ; à défaut, utiliser le tarif minute'],
                  ['Retour multi-jours', 'J+1 : 35 € ; J+2 : 70 € ; J+3 : 105 €', 'Même règle', 'Cumul ou montant unique à formaliser'],
              ], [1800, 1600, 1600, 4360], font_size=8.8)

    add_heading(doc, 'Retour à vide proposé', 2)
    add_table(doc,
              ['Distance de la course', 'Taux retour à vide', 'Calcul'],
              [
                  ['0 à 50 km', '0 %', 'Aucun montant de retour à vide'],
                  ['50 à 120 km', '30 %', 'Distance aller × tarif km × 30 %'],
                  ['120 à 180 km', '60 %', 'Distance aller × tarif km × 60 %'],
                  ['180 à 250 km', '70 %', 'Distance aller × tarif km × 70 %'],
                  ['Plus de 250 km', '80 % ou forfait négocié', 'Distance aller × tarif km × 80 %, sauf devis spécifique'],
              ], [3000, 2100, 4260], font_size=9)

    add_heading(doc, 'Atouts et contraintes', 2)
    add_bullet(doc, 'Atout : devis plus transparent, car chaque élément est nommé et additionné.')
    add_bullet(doc, 'Atout : le prix prend en compte le temps de conduite, utile en trafic dense ou en montagne.')
    add_bullet(doc, 'Atout : les coûts de retour sont encadrés par une règle commerciale simple.')
    add_bullet(doc, 'Contrainte : le prix dépend de la durée mesurée ou estimée ; il faut définir la source de vérité.')
    add_bullet(doc, 'Contrainte : la même course peut varier selon jour/nuit, durée, péage réel et statut multi-jours.')

    force_page_break(doc)

    # numerical comparison
    add_heading(doc, '3. Comparaison chiffrée sur les itinéraires exemples', 1)
    add_text(doc, 'Hypothèses communes : départ et prise en charge à Cluses ; aller simple ; tarif jour ; aucun péage, temps de trajet à 0 min, supplément à 0 € et remise à 0 €. La colonne « nouvelle formule » est donc une base hors temps.', 9.5, MUTED, italic=True, after=9)
    add_table(doc,
              ['Itinéraire', 'Distance', 'Ancienne formule', 'Nouvelle formule', 'Écart'],
              [
                  ['Cluses → Sallanches', '16 km', '42,24 €', '29,20 €', '−13,04 €'],
                  ['Cluses → Annecy', '69 km', '139,38 €', '117,64 €', '−21,74 €'],
                  ['Cluses → Lyon Saint-Exupéry', '130 km', '262,60 €', '259,60 €', '−3,00 €'],
                  ['Cluses → La Mure', '205 km', '414,10 €', '428,20 €', '+14,10 €'],
                  ['Cluses → Marseille', '477 km', '963,54 €', '1 040,32 €', '+76,78 €'],
              ], [3120, 950, 1800, 1900, 1590], font_size=9,
              aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT])

    add_heading(doc, 'Lecture des écarts', 2)
    add_bullet(doc, 'Sur les distances courtes et moyennes, le retour à vide forfaitisé peut être inférieur au retour CA facturé par l’ancienne formule.')
    add_bullet(doc, 'À partir de 180 km, le taux de retour à vide (70 % puis 80 %) peut rendre la nouvelle base plus élevée que l’ancienne formule.')
    add_bullet(doc, 'Dès qu’un temps de trajet est ajouté, la nouvelle formule augmente de 0,40 €/min jour ou 0,70 €/min nuit.')
    add_bullet(doc, 'Les chiffres ne sont pas comparables si les péages, la durée, les nuits ou les jours de retour sont différents.')

    add_heading(doc, 'Exemple détaillé : Cluses → Annecy (69 km, jour)', 2)
    add_table(doc,
              ['Ancienne formule publique', 'Nouvelle formule proposée (hors temps / péages)'],
              [
                  ['Trajet client : 69 km × 1,32 € = 91,08 €\nRetour CA : 69 km × 0,70 € = 48,30 €\nTotal : 139,38 €',
                   'Prise en charge : 10,00 €\nDistance : 69 km × 1,20 € = 82,80 €\nRetour à vide : 82,80 € × 30 % = 24,84 €\nBase : 117,64 €'],
              ], [4680, 4680], font_size=9)

    force_page_break(doc)

    # edge cases
    add_heading(doc, '4. Cas exceptionnels à traiter avant mise en production', 1)
    add_text(doc, 'Les règles ci-dessous doivent être écrites explicitement dans le calculateur, le devis et l’administration. Elles évitent les prix incohérents et les litiges.', 10.5, after=8)
    add_table(doc,
              ['Cas', 'Risque', 'Règle recommandée'],
              [
                  ['Course exactement à 50 / 120 / 180 / 250 km', 'Bascule de taux ambiguë.', 'Définir les bornes inclusives : par ex. 50,00 km reste à 0 %, puis 50,01 km passe à 30 %.'],
                  ['Distance > 250 km', 'Le « forfait négocié » peut être oublié.', 'Afficher une alerte et exiger une validation manuelle si aucun forfait n’est saisi.'],
                  ['Aller-retour client le même jour', 'Risque de facturer deux fois le retour à vide.', 'Définir une formule A/R distincte : TP ×2, puis décider s’il existe ou non un retour à vide.'],
                  ['Attente', 'Les 10 min peuvent être interprétées par arrêt ou par course.', 'Appliquer 10 min gratuites par course, puis arrondir selon une unité affichée (minute ou tranche).'],
                  ['Péages', 'Montant ×2 alors que le retour emprunte un autre itinéraire ou n’a pas de péage.', 'Enregistrer péage aller et péage retour séparément ; utiliser ×2 seulement si les deux sont identiques.'],
                  ['J+1 / J+2 / J+3', 'Cumul incertain et durée non contrôlée.', 'Préciser si 70 € remplace ou s’ajoute à 35 €, et bloquer au-delà de J+3 avec validation manuelle.'],
                  ['Nuit pendant la course', 'Départ jour, arrivée nuit.', 'Choisir : tarif selon heure de prise en charge (simple) ou découpage horaire (plus précis).'],
                  ['Remise', 'Remise > prix ou remise sur péages.', 'Interdire total négatif et définir l’assiette de remise (transport seul ou total TTC).'],
              ], [1950, 2760, 4650], font_size=8.6)

    add_heading(doc, 'Points techniques et comptables', 2)
    add_bullet(doc, 'TVA : fixer le taux applicable à chaque composant (transport, péage, attente, supplément) et conserver le détail HT/TVA/TTC dans la réservation.')
    add_bullet(doc, 'Arrondi : calculer avec précision interne, puis arrondir uniquement chaque total affiché à 2 décimales.')
    add_bullet(doc, 'Données d’itinéraire : figer distance, durée, péages et date/heure au moment de l’établissement du devis.')
    add_bullet(doc, 'Transparence : le client doit voir les lignes qui composent le total, notamment retour à vide, durée, péages et suppléments.')

    force_page_break(doc)

    # Implementation recommendation
    add_heading(doc, '5. Recommandation de mise en œuvre', 1)
    add_callout(doc, 'RECOMMANDATION', 'Adopter une seule fonction de calcul partagée par le calculateur public, les réservations, les devis et les factures. Les paramètres tarifaires doivent être administrables et versionnés.', PALE_GREEN)

    add_heading(doc, 'Règle de calcul proposée à formaliser', 2)
    add_table(doc,
              ['Étape', 'Calcul / contrôle'],
              [
                  ['1. Déterminer le tarif', 'Jour ou nuit selon la date et l’heure de prise en charge.'],
                  ['2. Fixer les données de route', 'Distance, durée, péage aller et péage retour issus de la même source d’itinéraire.'],
                  ['3. Calculer la base', '10 € + (distance × tarif km) + (durée × tarif minute).'],
                  ['4. Appliquer le retour à vide', 'Base distance uniquement × taux correspondant à la distance ; ne pas appliquer au temps ni aux péages.'],
                  ['5. Ajouter les frais', 'Péage aller + péage retour + attente + J+N.'],
                  ['6. Appliquer la remise', 'Selon une assiette et un plafond explicitement définis.'],
                  ['7. Produire le total', 'Arrondi TTC, détail HT/TVA/TTC, trace des paramètres utilisés.'],
              ], [2300, 7060], font_size=9)

    add_heading(doc, 'Checklist de validation avant déploiement', 2)
    add_bullet(doc, 'Valider la formule séparée pour A/S, A/R même jour et A/R sur plusieurs jours.')
    add_bullet(doc, 'Valider l’assiette du retour à vide : distance seule (recommandé) ou distance + temps.')
    add_bullet(doc, 'Valider la règle des péages : montants distincts aller/retour ou multiplication automatique.')
    add_bullet(doc, 'Valider le tarif d’attente après les 10 minutes gratuites.')
    add_bullet(doc, 'Valider la TVA de chaque ligne et les règles d’arrondi.')
    add_bullet(doc, 'Tester au minimum les bornes 50, 120, 180, 250 km, un dimanche, une course de nuit, un péage, une attente et un retour J+1.')

    add_text(doc, 'Conclusion : la nouvelle formule est commercialement plus simple et plus détaillable, mais elle doit être spécifiée précisément sur les cas A/R, péages, attente, TVA et forfaits longue distance avant de remplacer le moteur actuel.', 10.5, NAVY, True, before=12, after=0)

    OUT_DIR.mkdir(exist_ok=True)
    doc.save(DOCX_OUT)
    return DOCX_OUT


if __name__ == '__main__':
    path = build_document()
    print(path)
